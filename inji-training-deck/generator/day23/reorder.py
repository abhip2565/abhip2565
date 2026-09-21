from pptx import Presentation
import copy, sys

SRC, DST = 'out.pptx', 'Inji-Day2-Day3-Training-EXTENDED.pptx'
# new slides as built (1-indexed in out.pptx)
ANATOMY, TRUST, DEVICES, PAYLOADS, BINDING, BACKUP, QRLOGIN, SVG, BLE = range(26, 35)

ORDER = (list(range(1, 16))                 # 1-15  title .. FaceMatch
         + [SVG, BLE, QRLOGIN, BACKUP]      # new wallet-internals block
         + [16, 17, 18, 19, 20, 21]         # library summary .. auth request
         + [DEVICES, PAYLOADS]              # new OpenID4VP delivery + payloads
         + [22]                             # credential queries
         + [ANATOMY, TRUST]                 # new SD-JWT pair
         + [23]                             # VC vs VP
         + [BINDING]                        # new binding at presentation
         + [24, 25])                        # summary, status

prs = Presentation(SRC)
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
assert len(ORDER) == len(ids) == 34, (len(ORDER), len(ids))
assert sorted(ORDER) == list(range(1, 35))

for el in ids:
    sldIdLst.remove(el)
for pos in ORDER:
    sldIdLst.append(ids[pos - 1])

prs.save(DST)

# report
p2 = Presentation(DST)
NEW = {26, 27, 28, 29, 30, 31, 32, 33, 34}
for n, (pos, sl) in enumerate(zip(ORDER, p2.slides), 1):
    if pos not in NEW:
        continue
    t = ''
    for sh in sl.shapes:
        if sh.has_text_frame:
            tx = sh.text_frame.text.strip()
            if tx and not tx.isupper() and len(tx) > 10:
                t = tx.split('\n')[0]
                break
    print("  new slide %-3d %s" % (n, t[:56]))
print("saved", DST, "with", len(p2.slides._sldIdLst), "slides")
