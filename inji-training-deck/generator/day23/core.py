# Design system matched to the Day 2/3 Inji training deck (10 x 5.625 in).
import re, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

W, H = 10.0, 5.625
ML, MR = 0.62, 0.62
CW = W - ML - MR                    # 8.76
Y0 = 1.32                           # content top
YLIM = 5.20                         # content bottom

SANS  = 'Montserrat'
MED   = 'Montserrat Medium'
SEMI  = 'Montserrat SemiBold'
MONO  = 'Roboto Mono'

C = dict(
    white  = RGBColor(0xFF, 0xFF, 0xFF),
    body   = RGBColor(0xC7, 0xD0, 0xE8),
    body2  = RGBColor(0xE3, 0xE8, 0xF5),
    orange = RGBColor(0xF1, 0x75, 0x2F),
    orange2= RGBColor(0xF2, 0x7D, 0x21),
    rail   = RGBColor(0xF2, 0x68, 0x0C),
    mag    = RGBColor(0x95, 0x1F, 0x6F),
    mag2   = RGBColor(0xC0, 0x2D, 0x80),
    violet = RGBColor(0x8F, 0x35, 0xD6),
    violet2= RGBColor(0x6B, 0x71, 0xC1),
    indigo = RGBColor(0x3D, 0x25, 0x85),
    red    = RGBColor(0xD6, 0x31, 0x56),
    panel  = RGBColor(0x1D, 0x1E, 0x3A),
    dim    = RGBColor(0x8D, 0x98, 0xB8),
)

prs = None
BG = LOGO = None


def init(template, bg_png, logo_png):
    global prs, BG, LOGO
    prs = Presentation(template)
    BG, LOGO = bg_png, logo_png
    return prs


# ---------------------------------------------------------------- primitives
def _solid_alpha(shape, rgb, alpha_pct):
    """Fill with transparency (python-pptx has no alpha API)."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sf.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(alpha_pct * 1000))})
    clr.append(a)


def shadow_off(shape):
    sp = shape._element.spPr
    for tag in ('a:effectLst', 'a:effectRef'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    sp.append(sp.makeelement(qn('a:effectLst'), {}))


def rect(s, x, y, w, h, fill=None, alpha=None, line=None, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    elif alpha is not None:
        _solid_alpha(sh, fill, alpha)
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    shadow_off(sh)
    sh.text_frame.text = ''
    return sh


def tb(s, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return box, tf


def para(tf, text, size=11, color=None, bold=False, italic=False, font=None,
         align=PP_ALIGN.LEFT, first=False, space_before=0, line_spacing=1.2):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line_spacing
    if space_before:
        p.space_before = Pt(space_before)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font or SANS
    r.font.color.rgb = color if color is not None else C['body']
    return p


def rich(tf, chunks, size=11, align=PP_ALIGN.LEFT, first=False,
         line_spacing=1.2, space_before=0):
    """chunks: (text, color, bold[, italic[, font]])"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.line_spacing = line_spacing
    if space_before:
        p.space_before = Pt(space_before)
    for ch in chunks:
        txt, col, bold = ch[0], ch[1], ch[2]
        ital = ch[3] if len(ch) > 3 else False
        fnt = ch[4] if len(ch) > 4 else SANS
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = ital
        r.font.name = fnt; r.font.color.rgb = col
    return p


def md(tf, text, size=11, color=None, align=PP_ALIGN.LEFT, first=False,
       line_spacing=1.2, code_color=None, space_before=0, base_bold=False):
    """Inline `code` and **bold**."""
    base = color if color is not None else C['body']
    cc = code_color or C['orange']
    out = []
    for pt in re.split(r'(`[^`]+`|\*\*[^*`]+\*\*)', text):
        if not pt:
            continue
        if pt.startswith('`'):
            out.append((pt[1:-1], cc, base_bold, False, MONO))
        elif pt.startswith('**'):
            out.append((pt[2:-2], C['white'], True))
        else:
            out.append((pt, base, base_bold))
    return rich(tf, out, size=size, align=align, first=first,
                line_spacing=line_spacing, space_before=space_before)


def line(s, x1, y1, x2, y2, color=None, lw=1.25, dash=None):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color or C['dim']
    c.line.width = Pt(lw)
    if dash:
        ln = c.line._get_or_add_ln()
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'}); ln.append(d)
    c.shadow.inherit = False
    return c


def arrowhead(shape):
    ln = shape.line._get_or_add_ln()
    e = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(e)


def arrow(s, x1, y1, x2, y2, color=None, lw=1.3, dash=None):
    c = line(s, x1, y1, x2, y2, color, lw, dash)
    arrowhead(c)
    return c


def chip(s, x, y, w, h, text, fill, fg=None, size=8, bold=True):
    sh = rect(s, x, y, w, h, fill=fill, adj=0.35)
    tf = sh.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = SANS
    r.font.color.rgb = fg or C['white']
    return sh


# ---------------------------------------------------------------- slide chrome
def new_slide():
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    s = prs.slides.add_slide(blank)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s.shapes.add_picture(BG, 0, 0, Inches(W), Inches(H))
    return s


def head(s, title, kicker):
    line(s, 0.47, 0.42, 0.47, 5.11, C['rail'], 2.25)
    _, tf = tb(s, ML, 0.42, CW, 0.31)
    para(tf, kicker.upper(), size=10, color=C['body'], bold=True, first=True)
    _, tf = tb(s, ML, 0.67, CW, 0.52)
    para(tf, title, size=23, color=C['white'], bold=True, italic=True, first=True)
    s.shapes.add_picture(LOGO, Inches(9.02), Inches(0.26), Inches(0.67), Inches(0.36))
    return Y0


def notes(s, lines, minutes=None):
    tf = s.notes_slide.notes_text_frame
    body = "\n\n".join(lines)
    if minutes:
        body += "\n\n[%s]" % minutes
    tf.text = body


# ---------------------------------------------------------------- blocks
def panel(s, x, y, w, h, accent=None, alpha=5.0):
    sh = rect(s, x, y, w, h, fill=C['white'], alpha=alpha,
              line=accent, lw=1.0 if accent else 0)
    return sh


def card(s, x, y, w, h, title, items, accent, tsize=10.5, isize=8.6, alpha=5.0):
    panel(s, x, y, w, h, accent, alpha)
    rect(s, x, y, w, 0.045, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    _, tf = tb(s, x + 0.14, y + 0.16, w - 0.28, 0.26)
    para(tf, title, size=tsize, color=accent, bold=True, first=True)
    cy = y + 0.46
    for it in items:
        rect(s, x + 0.16, cy + 0.055, 0.05, 0.05, fill=accent, shape=MSO_SHAPE.OVAL)
        _, t2 = tb(s, x + 0.30, cy, w - 0.46, 0.5)
        md(t2, it, size=isize, first=True, line_spacing=1.2)
        cap = max(12, int((w - 0.46) * (72.0 / (isize * 0.52))))
        nl = max(1, -(-len(re.sub(r'[`*]', '', it)) // cap))
        cy += nl * (isize / 72.0) * 1.25 + 0.085
    return y + h


def codebox(s, x, y, w, h, lines, label=None, size=7.4, hl=(), fill_alpha=9.0):
    if label:
        _, tf = tb(s, x, y - 0.21, w, 0.2)
        para(tf, label, size=7.6, color=C['dim'], bold=True, first=True)
    rect(s, x, y, w, h, fill=C['panel'], alpha=None, adj=0.04)
    _, tf = tb(s, x + 0.12, y + 0.10, w - 0.24, h - 0.20)
    for i, ln in enumerate(lines):
        col = C['orange'] if i in hl else C['body']
        if ln.startswith('//') or ln.startswith('#'):
            col = C['dim']
        para(tf, ln if ln else ' ', size=size, color=col, font=MONO,
             first=(i == 0), line_spacing=1.22)
    return y + h


def callout(s, x, y, w, text, accent=None, size=9.0, h=None, title=None):
    accent = accent or C['orange']
    if h is None:
        cap = int((w - 0.5) * (72.0 / (size * 0.50)) * 0.96)
        nl = max(1, -(-len(re.sub(r'[`*]', '', text)) // max(20, cap)))
        h = 0.18 + nl * (size / 72.0) * 1.34 + (0.20 if title else 0)
    rect(s, x, y, w, h, fill=accent, alpha=14.0, adj=0.10)
    rect(s, x, y, 0.04, h, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    _, tf = tb(s, x + 0.18, y + 0.085, w - 0.32, h - 0.15)
    if title:
        para(tf, title, size=size - 0.6, color=accent, bold=True, first=True)
        md(tf, text, size=size, code_color=accent, line_spacing=1.28)
    else:
        md(tf, text, size=size, code_color=accent, first=True, line_spacing=1.28)
    return y + h


def rowtable(s, x, y, w, headers, rows, col_w, fsize=8.4, hsize=8.6,
             accent=None, label_col=True):
    """Matches the deck's label + tinted-cell row style."""
    accent = accent or C['mag']
    tot = float(sum(col_w))
    xs, acc = [], x
    for cwv in col_w:
        xs.append(acc); acc += w * (cwv / tot)
    ws = [w * (cwv / tot) for cwv in col_w]
    # header
    for j, hdr in enumerate(headers):
        if j == 0 and label_col:
            continue
        if not hdr:
            continue
        fill = accent if j == len(headers) - 1 else C['white']
        alpha = None if j == len(headers) - 1 else 7.8
        rect(s, xs[j], y, ws[j] - 0.06, 0.30, fill=fill, alpha=alpha, adj=0.14)
        _, tf = tb(s, xs[j] + 0.08, y, ws[j] - 0.22, 0.30, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, hdr, size=hsize, color=C['white'] if alpha is None else C['body'],
             bold=True, align=PP_ALIGN.CENTER, first=True)
    yy = y + 0.38
    for r in rows:
        nl = 1
        for j, v in enumerate(r):
            if j == 0:
                continue
            cap = max(10, int((ws[j] - 0.3) * (72.0 / (fsize * 0.50))))
            nl = max(nl, -(-len(re.sub(r'[`*]', '', str(v))) // cap))
        rh = max(0.34, 0.14 + nl * (fsize / 72.0) * 1.30)
        for j, v in enumerate(r):
            if j == 0:
                _, tf = tb(s, xs[j], yy, ws[j] - 0.10, rh, anchor=MSO_ANCHOR.MIDDLE)
                md(tf, v, size=fsize + 0.5, color=C['white'], first=True)
                continue
            last = (j == len(r) - 1)
            rect(s, xs[j], yy, ws[j] - 0.06, rh, fill=C['white'],
                 alpha=9.0 if last else 4.2,
                 line=accent if last else None, lw=0.75, adj=0.10)
            _, tf = tb(s, xs[j] + 0.11, yy + 0.06, ws[j] - 0.28, rh - 0.12)
            md(tf, str(v), size=fsize, color=C['white'] if last else C['body'],
               first=True, line_spacing=1.24)
        yy += rh + 0.08
    return yy


def steps(s, y, items, h=0.98, gap=0.12, nsize=7.2, tsize=9.4, ssize=7.4):
    n = len(items)
    bw = (CW - gap * (n - 1)) / n
    for i, (num, title, sub, col) in enumerate(items):
        x = ML + i * (bw + gap)
        panel(s, x, y, bw, h, col, 6.0)
        _, tf = tb(s, x + 0.11, y + 0.08, bw - 0.22, 0.16)
        para(tf, num, size=nsize, color=col, bold=True, first=True)
        _, tf = tb(s, x + 0.11, y + 0.26, bw - 0.22, 0.32)
        md(tf, title, size=tsize, color=C['white'], base_bold=True,
           code_color=col, first=True, line_spacing=1.06)
        _, tf = tb(s, x + 0.11, y + 0.58, bw - 0.22, h - 0.64)
        md(tf, sub, size=ssize, color=C['body'], code_color=col,
           first=True, line_spacing=1.14)
        if i < n - 1:
            arrow(s, x + bw + 0.012, y + h / 2, x + bw + gap - 0.012, y + h / 2,
                  C['dim'], 1.0)
    return y + h


# ---------------------------------------------------------------- checks
def bounds_report():
    bad = []
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
            except TypeError:
                continue
            if None in (l, t, w, h):
                continue
            if h / 914400 > 5.0:
                continue
            L, R, B = l / 914400, (l + w) / 914400, (t + h) / 914400
            txt = sh.text_frame.text.strip()[:30] if sh.has_text_frame else ''
            if L < -0.02 or R > W + 0.02:
                bad.append((i, 'HORIZ', round(L, 2), round(R, 2), txt))
            if B > 5.30:
                bad.append((i, 'BOTTOM', round(B, 2), '', txt))
    return bad
