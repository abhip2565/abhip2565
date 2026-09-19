from pptx import Presentation
p = Presentation('preview.pptx')
EMU = 914400
bad = []
for i, sl in enumerate(p.slides, 1):
    worst = 0; title = ''
    for sh in sl.shapes:
        try:
            t, h, l, w = sh.top, sh.height, sh.left, sh.width
        except TypeError:
            continue
        if t is None: continue
        bi=(t+h)/EMU; li=l/EMU; wi=w/EMU; hi=h/EMU
        if hi > 6.9: continue
        if abs(t/EMU-7.02)<0.01 and hi<0.35: continue
        if sh.has_text_frame and not title:
            tx = sh.text_frame.text.strip().split('\n')[0]
            if len(tx) > 12: title = tx[:52]
        if bi > worst: worst = bi
    if worst > 6.90:
        bad.append((i, round(worst,2), title))
for b in bad: print("%-4d %.2f  %s" % b)
print(len(bad), "to fix")
