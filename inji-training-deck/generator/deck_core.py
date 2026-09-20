# -*- coding: utf-8 -*-
"""Design system + primitive builders for the Inji technical training deck."""
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- tokens
C = dict(
    ink=RGBColor(0x0F, 0x17, 0x2A),
    ink2=RGBColor(0x1E, 0x29, 0x3B),
    muted=RGBColor(0x5B, 0x6B, 0x7F),
    faint=RGBColor(0x94, 0xA3, 0xB8),
    line=RGBColor(0xCB, 0xD5, 0xE1),
    surf=RGBColor(0xF3, 0xF6, 0xF9),
    surf2=RGBColor(0xE4, 0xEA, 0xF1),
    white=RGBColor(0xFF, 0xFF, 0xFF),
    primary=RGBColor(0x1D, 0x4E, 0x89),
    primary_d=RGBColor(0x12, 0x30, 0x4F),
    primary_l=RGBColor(0xDC, 0xE7, 0xF4),
    accent=RGBColor(0x0E, 0x94, 0x88),
    accent_l=RGBColor(0xD6, 0xF0, 0xED),
    amber=RGBColor(0xB4, 0x53, 0x09),
    amber_l=RGBColor(0xFD, 0xF0, 0xDC),
    red=RGBColor(0xB3, 0x1B, 0x1B),
    red_l=RGBColor(0xFC, 0xE6, 0xE6),
    violet=RGBColor(0x6D, 0x28, 0xD9),
    violet_l=RGBColor(0xEE, 0xE8, 0xFD),
    green=RGBColor(0x15, 0x70, 0x3D),
    green_l=RGBColor(0xDF, 0xF1, 0xE5),
    dark=RGBColor(0x0B, 0x1F, 0x33),
    dark2=RGBColor(0x14, 0x2E, 0x49),
    code_bg=RGBColor(0x0E, 0x1C, 0x2D),
    bg=RGBColor(0xFF, 0xFF, 0xFF),
)
F_SANS = "Arial"
F_MONO = "Consolas"

W, H = 13.333, 7.5
ML, MR = 0.55, 0.55
CW = W - ML - MR          # content width 12.233
BODY_TOP = 1.34
BODY_BOT = 6.88

THEME = dict(
    bg_image=None, logo=None, rail=None, mark=None,
    canvas=False, canvas_fill=None, canvas_top=None,
    title_color=None, sub_color=None, kicker_color=None, footer_color=None,
    sub_code_color=None, sub_strong_color=None,
    title_italic=False, rule=True,
)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

STATE = dict(section="", n=0, toc=[])

# ---------------------------------------------------------------- xml helpers
def _spPr(shape):
    return shape._element.spPr

def _get_ln(shape):
    sp = _spPr(shape)
    ln = sp.find(qn('a:ln'))
    if ln is None:
        ln = sp.makeelement(qn('a:ln'), {})
        sp.append(ln)
    return ln

def arrowhead(shape, tail=True, head=False, typ='triangle', size='med'):
    ln = _get_ln(shape)
    for on, tag in ((tail, 'a:tailEnd'), (head, 'a:headEnd')):
        if not on:
            continue
        e = ln.find(qn(tag))
        if e is None:
            e = ln.makeelement(qn(tag), {})
            ln.append(e)
        e.set('type', typ); e.set('w', size); e.set('len', size)

def no_line(shape):
    shape.line.fill.background()

def shadow_off(shape):
    sp = _spPr(shape)
    for t in ('a:effectLst',):
        if sp.find(qn(t)) is None:
            sp.append(sp.makeelement(qn(t), {}))

# ---------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         dash=None, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if adj is not None:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        no_line(s)
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    s.shadow.inherit = False
    return s

def tb(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return t, tf

def para(tf, text, size=12, color=None, bold=False, italic=False, font=None,
         align=PP_ALIGN.LEFT, space_before=0, space_after=0, first=False,
         line_spacing=1.14):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run(); r.text = text
    f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font
    f.color.rgb = color if color is not None else C['ink']
    return p

def rich(tf, chunks, size=12, align=PP_ALIGN.LEFT, first=False, space_before=0,
         space_after=0, font=None, line_spacing=1.14):
    """chunks = [(text, color, bold, italic?, font?), ...]"""
    font = font or F_SANS
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before); p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for ch in chunks:
        txt = ch[0]; col = ch[1] if len(ch) > 1 and ch[1] is not None else C['ink']
        bold = ch[2] if len(ch) > 2 else False
        ital = ch[3] if len(ch) > 3 else False
        fnt = ch[4] if len(ch) > 4 else font
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = ital
        r.font.name = fnt; r.font.color.rgb = col
    return p

def line(slide, x1, y1, x2, y2, color=None, lw=1.25, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color if color is not None else C['line']
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    c.shadow.inherit = False
    return c

def chip(slide, x, y, w, h, text, fill, fg, size=9.5, bold=True, border=None):
    s = rect(slide, x, y, w, h, fill=fill, line=border, lw=0.9,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.22)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = F_SANS
    r.font.color.rgb = fg
    return s

# ---------------------------------------------------------------- chrome
def new_slide(dark=False, chrome=True):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = C['dark'] if dark else C['bg']
    if THEME.get('bg_image'):
        s.shapes.add_picture(THEME['bg_image'], 0, 0, Inches(W), Inches(H))
        if chrome and THEME.get('rail'):
            r = s.shapes.add_picture(THEME['rail'], Inches(-3.02), Inches(3.68),
                                     Inches(6.40), Inches(0.09))
            r.rotation = 90
        if chrome and THEME.get('logo'):
            s.shapes.add_picture(THEME['logo'], Inches(12.00), Inches(0.30),
                                 Inches(0.89), Inches(0.45))
    STATE['n'] += 1
    return s

def notes(slide, explain, caveats=None, questions=None, minutes=None):
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "WHAT TO EXPLAIN"
    r.font.bold = True; r.font.size = Pt(11)
    for ln_ in explain if isinstance(explain, (list, tuple)) else [explain]:
        q = tf.add_paragraph(); rr = q.add_run(); rr.text = "• " + ln_
        rr.font.size = Pt(11)
    if caveats:
        q = tf.add_paragraph(); rr = q.add_run(); rr.text = "CAVEATS / GOTCHAS"
        rr.font.bold = True; rr.font.size = Pt(11)
        for ln_ in caveats:
            q = tf.add_paragraph(); rr = q.add_run(); rr.text = "• " + ln_
            rr.font.size = Pt(11)
    if questions:
        q = tf.add_paragraph(); rr = q.add_run(); rr.text = "LIKELY QUESTIONS"
        rr.font.bold = True; rr.font.size = Pt(11)
        for ln_ in questions:
            q = tf.add_paragraph(); rr = q.add_run(); rr.text = "– " + ln_
            rr.font.size = Pt(11)
    if minutes:
        q = tf.add_paragraph(); rr = q.add_run(); rr.text = "TIME: %s" % minutes
        rr.font.bold = True; rr.font.size = Pt(11)

def footer(slide, dark=False):
    fg = THEME.get('footer_color') or (C['faint'] if not dark
                                       else RGBColor(0x64, 0x80, 0x9B))
    _, tf = tb(slide, ML, 7.02, 9.0, 0.3)
    para(tf, STATE['section'], size=8.5, color=fg, first=True)
    _, tf2 = tb(slide, W - MR - 1.2, 7.02, 1.2, 0.3)
    para(tf2, str(STATE['n']), size=8.5, color=fg, align=PP_ALIGN.RIGHT, first=True)

def head(slide, title, kicker=None, sub=None, rule=True):
    y = 0.42
    if kicker:
        _, tf = tb(slide, ML, y - 0.07, CW, 0.26)
        para(tf, kicker.upper(), size=9,
             color=THEME.get('kicker_color') or C['accent'], bold=True, first=True)
        y += 0.28
    _, tf = tb(slide, ML, y, CW, 0.52)
    para(tf, title, size=25, color=THEME.get('title_color') or C['ink'],
         bold=True, italic=bool(THEME.get('title_italic')), first=True, line_spacing=1.0)
    y2 = y + (0.50 if not sub else 0.48)
    if sub:
        _, tf = tb(slide, ML, y2, CW, 0.34)
        ch = []
        for pt in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', sub):
            if not pt:
                continue
            if pt.startswith('`'):
                ch.append((pt[1:-1], THEME.get('sub_code_color') or C['primary'],
                           False, False, F_MONO))
            elif pt.startswith('**'):
                ch.append((pt[2:-2], THEME.get('sub_strong_color') or C['ink'], True))
            else:
                ch.append((pt, THEME.get('sub_color') or C['muted'], False))
        rich(tf, ch, size=11.5, first=True)
        y2 += 0.34
    if rule and THEME.get('rule', True):
        line(slide, ML, y2 + 0.10, W - MR, y2 + 0.10, C['line'], 1.0)
    return y2 + 0.32

# ---------------------------------------------------------------- code box
KEY_RE = re.compile(r'("(?:[^"\\]|\\.)*")(\s*:)')
STR_RE = re.compile(r'"(?:[^"\\]|\\.)*"')

def _tokenise(text):
    """crude JSON/HTTP highlighter -> list of (chunk, colorkey)"""
    out = []
    i = 0
    s = text
    if s.strip().startswith(('#', '//')):
        return [(s, 'cmt')]
    # split by comment marker at end
    cmt = None
    m = re.search(r'\s+(//.*|#\s.*)$', s)
    if m:
        cmt = m.group(1); s = s[:m.start()]
    spans = []
    for m in STR_RE.finditer(s):
        spans.append((m.start(), m.end()))
    pos = 0
    for a, b in spans:
        if a > pos:
            out.append((s[pos:a], 'plain'))
        lit = s[a:b]
        after = s[b:b + 2]
        out.append((lit, 'key' if after.lstrip().startswith(':') else 'str'))
        pos = b
    if pos < len(s):
        out.append((s[pos:], 'plain'))
    if cmt:
        out.append(("   " + cmt.strip(), 'cmt'))
    return out

CODE_COL = dict(
    plain=RGBColor(0xCB, 0xD9, 0xE8),
    key=RGBColor(0x7D, 0xD3, 0xC8),
    str=RGBColor(0xF6, 0xC6, 0x7A),
    cmt=RGBColor(0x7E, 0x93, 0xA8),
)

def codebox(slide, x, y, w, h, lines, label=None, size=9.2, hl=None):
    hl = hl or []
    if label:
        _, tf = tb(slide, x, y - 0.24, w, 0.22)
        para(tf, label, size=9, color=C['muted'], bold=True, first=True)
    box = rect(slide, x, y, w, h, fill=C['code_bg'], line=None,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.04)
    t, tf = tb(slide, x + 0.16, y + 0.13, w - 0.32, h - 0.26)
    for i, ln_ in enumerate(lines):
        chunks = _tokenise(ln_)
        col_over = RGBColor(0xFF, 0xD5, 0x6B) if i in hl else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.16
        for chunk, kind in chunks:
            r = p.add_run(); r.text = chunk
            r.font.size = Pt(size); r.font.name = F_MONO
            r.font.color.rgb = col_over or CODE_COL[kind]
            if i in hl:
                r.font.bold = True
    return box

# ---------------------------------------------------------------- tables
def table(slide, x, y, w, headers, rows, col_w=None, fsize=9.5, hsize=9.5,
          row_h=0.3, head_h=0.34, zebra=True, accent=None, align=None):
    accent = accent or C['primary']
    ncols = len(headers); nrows = len(rows) + 1
    # estimate wrapped height per row so callers can place content underneath
    widths = col_w or [1.0] * ncols
    tot_w = float(sum(widths))
    cw_in = [w * (cwi / tot_w) for cwi in widths]
    def _lines(txt, cwi, fs):
        plain = re.sub(r'[`*+!~]', '', str(txt))
        if not plain:
            return 1
        cap = max(6, int((cwi - 0.15) * (72.0 / (fs * 0.505)) * 0.95))
        return max(1, -(-len(plain) // cap))
    head_hh = max(head_h, 0.075 + max(_lines(h, cw_in[j], hsize)
                                      for j, h in enumerate(headers)) * hsize * 1.30 / 72.0)
    row_hh = []
    for r in rows:
        nl = max(_lines(v, cw_in[j], fsize) for j, v in enumerate(r))
        row_hh.append(max(row_h, 0.065 + nl * fsize * 1.30 / 72.0))
    total_h = head_hh + sum(row_hh)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w),
                                Inches(total_h))
    t = gf.table
    tblPr = t._tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        for el in tblPr.findall(qn('a:tableStyleId')):
            tblPr.remove(el)
        sid = tblPr.makeelement(qn('a:tableStyleId'), {})
        sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
        tblPr.append(sid)
    t.first_row = True
    t.horz_banding = False
    if col_w:
        tot = float(sum(col_w))
        for i, cwi in enumerate(col_w):
            t.columns[i].width = Emu(int(Inches(w) * (cwi / tot)))
    t.rows[0].height = Inches(head_hh)
    for i in range(len(rows)):
        t.rows[i + 1].height = Inches(row_hh[i])
    for j, htxt in enumerate(headers):
        c = t.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = accent
        c.margin_left = Inches(0.07); c.margin_right = Inches(0.06)
        c.margin_top = Inches(0.03); c.margin_bottom = Inches(0.03)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = (align[j] if align else PP_ALIGN.LEFT)
        for seg in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', str(htxt)):
            if not seg:
                continue
            fnt = F_SANS
            if seg.startswith('`') and seg.endswith('`'):
                seg = seg[1:-1]; fnt = F_MONO
            elif seg.startswith('**') and seg.endswith('**'):
                seg = seg[2:-2]
            r = p.add_run(); r.text = seg
            r.font.size = Pt(hsize); r.font.bold = True; r.font.name = fnt
            r.font.color.rgb = C['white']
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i + 1, j)
            c.fill.solid()
            c.fill.fore_color.rgb = C['bg'] if (not zebra or i % 2 == 0) else C['surf']
            c.margin_left = Inches(0.07); c.margin_right = Inches(0.06)
            c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = (align[j] if align else PP_ALIGN.LEFT)
            p.line_spacing = 1.05
            txt = str(val); base_bold = False; base_col = C['ink2']
            if txt.startswith('!!'):
                txt = txt[2:]; base_col = C['red']; base_bold = True
            elif txt.startswith('++'):
                txt = txt[2:]; base_col = C['green']; base_bold = True
            elif txt.startswith('~~'):
                txt = txt[2:]; base_col = C['amber']; base_bold = True
            for seg in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', txt):
                if not seg:
                    continue
                fnt = F_SANS; bold = base_bold; col = base_col
                if seg.startswith('**') and seg.endswith('**'):
                    seg = seg[2:-2]; bold = True
                    col = base_col if base_col is not C['ink2'] else C['ink']
                elif seg.startswith('`') and seg.endswith('`'):
                    seg = seg[1:-1]; fnt = F_MONO
                    col = base_col if base_col is not C['ink2'] else C['primary']
                r = p.add_run(); r.text = seg
                r.font.size = Pt(fsize); r.font.bold = bold; r.font.name = fnt
                r.font.color.rgb = col
    return y + total_h


# ---------------------------------------------------------------- auto-fit
def _runs(shape):
    if getattr(shape, 'has_text_frame', False):
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                yield r

def fit_slide(sl, top0=1.28, limit=6.86, footer_top=7.0):
    EMU = 914400
    movable = []
    worst = 0.0
    for sh in sl.shapes:
        try:
            t, h, w = sh.top, sh.height, sh.width
        except TypeError:
            continue
        if t is None or h is None:
            continue
        if sh.name == 'BODYCANVAS':
            continue                      # stretched to the body after fitting
        ti, hi, wi = t / EMU, h / EMU, w / EMU
        if hi > 6.9 or (wi > 13.0 and hi > 7.0):
            continue                      # backgrounds / full-height rails
        if ti >= footer_top - 0.02 and hi < 0.35:
            continue                      # footer
        if ti + hi <= top0:
            continue                      # header chrome
        movable.append(sh)
        worst = max(worst, ti + hi)
    if worst <= limit or not movable:
        return 0.0
    k = (limit - top0) / (worst - top0)
    if k < 0.70:
        k = 0.70
    for sh in movable:
        t = sh.top / EMU
        nt = top0 + (t - top0) * k if t > top0 else t
        sh.top = Emu(int(nt * EMU))
        sh.height = Emu(int(sh.height * k))
        if sh.has_text_frame:
            for r in _runs(sh):
                if r.font.size is not None:
                    r.font.size = Pt(max(6.8, round(r.font.size.pt * (0.55 + 0.45 * k), 1)))
        if getattr(sh, 'has_table', False):
            tbl = sh.table
            for row in tbl.rows:
                row.height = Emu(int(row.height * k))
            for row in tbl.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size is not None:
                                r.font.size = Pt(max(6.8, round(r.font.size.pt * (0.55 + 0.45 * k), 1)))
    return k

def _stretch_canvas(sl, bottom=6.94):
    for sh in sl.shapes:
        if sh.name == 'BODYCANVAS':
            sh.height = Emu(int((bottom - sh.top / 914400) * 914400))


def fit_all(canvas_bottom=6.94):
    n = 0
    for sl in prs.slides:
        if fit_slide(sl):
            n += 1
        _stretch_canvas(sl, canvas_bottom)
    return n
